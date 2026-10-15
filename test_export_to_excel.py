import csv
import datetime as dt
import gzip
import os
import uuid
from decimal import Decimal
from pathlib import Path

import psycopg2
import pytest
from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

import export_to_excel as ex

DB = {
    "host": os.environ.get("EXPORT_TEST_HOST", "127.0.0.1"),
    "port": int(os.environ.get("EXPORT_TEST_PORT", "5432")),
    "dbname": os.environ.get("EXPORT_TEST_DB", "hospital_export_test"),
    "user": os.environ.get("EXPORT_TEST_USER", "postgres"),
}
DB_ARGS = ["--host", DB["host"], "--port", str(DB["port"]),
           "--dbname", DB["dbname"], "--user", DB["user"]]

pytestmark = pytest.mark.skipif(
    not os.environ.get("PGPASSWORD"),
    reason="set PGPASSWORD to run the tests against PostgreSQL",
)

LONG_DESCRIPTION = ("Patient history line with details. " * 40 + "\n") * 4


def connect():
    return psycopg2.connect(password=os.environ["PGPASSWORD"], **DB)


def query_rows(query: str) -> list[tuple]:
    with connect() as connection, connection.cursor() as cursor:
        cursor.execute(query)
        return cursor.fetchall()


def run_export(tmp_path: Path, output_name: str, query: str | None = None, *extra) -> tuple[int, Path]:
    output = tmp_path / "out" / output_name
    argv = ["--output", str(output), "--progress-seconds", "1000", *DB_ARGS, *extra]
    if query is not None:
        query_file = tmp_path / "query.sql"
        query_file.write_text(query, encoding="utf-8")
        argv += ["--query-file", str(query_file)]
    return ex.main(argv), output


def excel_expected(value):
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value if len(str(abs(value))) <= 15 else str(value)
    if isinstance(value, Decimal):
        return float(value) if len(value.as_tuple().digits) <= 15 else str(value)
    if isinstance(value, dt.datetime):
        if value.tzinfo is not None:
            value = value.astimezone(dt.timezone.utc).replace(tzinfo=None)
        return value if value.date() >= dt.date(1900, 3, 1) else value.isoformat(" ")
    if isinstance(value, dt.date):
        if value >= dt.date(1900, 3, 1):
            return dt.datetime.combine(value, dt.time())
        return value.isoformat()
    text = str(value)
    return ILLEGAL_CHARACTERS_RE.sub(lambda m: f"\\x{ord(m.group()):02x}", text)


def assert_same(read, expected, where):
    if isinstance(expected, float):
        assert isinstance(read, (int, float)), where
        assert abs(read - expected) <= 1e-9 * max(1.0, abs(expected)), where
    elif expected == "":
        assert read in ("", None), where
    else:
        assert read == expected, where


def data_rows(path: Path) -> list[tuple]:
    workbook = load_workbook(path, read_only=True)
    rows = []
    for sheet in workbook.worksheets:
        rows += list(sheet.iter_rows(min_row=2, values_only=True))
    workbook.close()
    return rows


def assert_matches_query(rows_read, db_rows):
    assert len(rows_read) == len(db_rows)
    for number, (read, source) in enumerate(zip(rows_read, db_rows), start=1):
        assert len(read) <= len(source), f"row {number} has extra columns"
        read = tuple(read) + (None,) * (len(source) - len(read))
        for column, (r, s) in enumerate(zip(read, source)):
            assert_same(r, excel_expected(s), f"row {number}, column {column + 1}")


@pytest.fixture
def edge_schema():
    schema = f"export_test_{uuid.uuid4().hex[:8]}"
    with connect() as connection, connection.cursor() as cursor:
        cursor.execute(f"""
            CREATE SCHEMA {schema};
            CREATE TABLE {schema}.doctor (
                doctor_id int PRIMARY KEY, doctor_name text NOT NULL, fee numeric(10, 2));
            CREATE TABLE {schema}.patient_details (
                patient_id int PRIMARY KEY, patient_name text NOT NULL, description text,
                admitted_on date, discharged_at timestamptz, balance numeric(12, 2),
                big_number numeric, visits bigint,
                doctor_id int REFERENCES {schema}.doctor (doctor_id));
        """)
        cursor.executemany(
            f"INSERT INTO {schema}.doctor VALUES (%s, %s, %s)",
            [(1, "Dr. Ayesha Khan", Decimal("1500.50")),
             (2, '=HYPERLINK("http://example.com")', None)],
        )
        cursor.executemany(
            f"INSERT INTO {schema}.patient_details VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)",
            [
                (1, "Ali Raza", LONG_DESCRIPTION, dt.date(2024, 1, 15),
                 "2024-01-20 10:30:00+05", Decimal("1234.56"),
                 Decimal("12345678901234567890.123"), 12, 1),
                (2, "=SUM(A1:A2)", "", None, None, None, None, None, None),
                (3, "Müller 患者 😀", 'Line 1\nLine 2, "quoted", comma; semi', dt.date(1899, 12, 31),
                 None, Decimal("-0.01"), None, 10**17, 2),
                (4, "#N/A", "bell\x07char", dt.date(2025, 12, 31), None, Decimal("0"), None, 0, 1),
                (5, "Zara", None, None, None, None, None, None, None),
            ],
        )
    yield schema
    with connect() as connection, connection.cursor() as cursor:
        cursor.execute(f"DROP SCHEMA {schema} CASCADE")


def edge_query(schema: str) -> str:
    return f"""
        SELECT p.patient_id, p.patient_name, p.description, p.admitted_on,
               p.discharged_at, p.balance, p.big_number, p.visits,
               d.doctor_id, d.doctor_name, d.fee AS doctor_fee
        FROM {schema}.patient_details AS p
        LEFT JOIN {schema}.doctor AS d ON d.doctor_id = p.doctor_id
        ORDER BY p.patient_id
    """


def test_patient_doctor_join_matches_the_database(tmp_path):
    query = ex.DEFAULT_QUERY_FILE.read_text(encoding="utf-8")
    code, output = run_export(tmp_path, "patients.xlsx")

    assert code == 0
    source = query_rows(query)
    rows = data_rows(output)
    assert_matches_query(rows, source)

    without_doctor = query_rows(
        "SELECT count(*) FROM patient_details p "
        "LEFT JOIN doctor d ON d.doctor_id = p.doctor_id WHERE d.doctor_id IS NULL"
    )[0][0]
    exported_without_doctor = sum(1 for row in rows if len(row) < 8 or row[7] is None)
    assert exported_without_doctor == without_doctor
    print(f"\n{len(rows):,} joined rows exported; {without_doctor:,} patients have no doctor")

    workbook = load_workbook(output)
    sheet = workbook["Data"]
    assert sheet.max_column == 10
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == f"A1:J{len(rows) + 1}"
    headers = [cell.value for cell in sheet[1]]
    assert headers[:3] == ["Patient ID", "Patient Name", "Patient Age"]
    assert sheet["F2"].number_format == "yyyy-mm-dd"
    assert all(sheet.column_dimensions[c].width <= 60 for c in "ABCDEFGHIJ")
    assert not list(output.parent.glob("*.part"))


def test_edge_cases_are_written_exactly(tmp_path, edge_schema):
    query = edge_query(edge_schema)
    code, output = run_export(tmp_path, "edge.xlsx", query)

    assert code == 0
    assert_matches_query(data_rows(output), query_rows(query))

    sheet = load_workbook(output)["Data"]
    assert sheet.max_column == 11
    assert sheet.column_dimensions["C"].width == 60
    assert sheet["C2"].alignment.wrap_text is True
    assert sheet["C2"].value == LONG_DESCRIPTION
    assert sheet.row_dimensions[2].height == 150
    assert (sheet["B3"].value, sheet["B3"].data_type) == ("=SUM(A1:A2)", "s")
    assert (sheet["B5"].value, sheet["B5"].data_type) == ("#N/A", "s")
    assert sheet["J3"].value is None
    assert sheet["J4"].data_type == "s"
    assert sheet["B4"].value == "Müller 患者 😀"
    assert sheet["C4"].value.startswith("Line 1\nLine 2")
    assert sheet["C5"].value == "bell\\x07char"
    assert sheet["D2"].value == dt.datetime(2024, 1, 15)
    assert sheet["D2"].number_format == "yyyy-mm-dd"
    assert sheet["D4"].value == "1899-12-31"
    assert sheet["E2"].value == dt.datetime(2024, 1, 20, 5, 30)
    assert sheet["E1"].value == "Discharged At (UTC)"
    assert sheet["F2"].number_format == "#,##0.00"
    assert sheet["G2"].value == "12345678901234567890.123"
    assert sheet["H4"].value == str(10**17)
    assert sheet["C6"].value is None and sheet["K6"].value is None


def test_csv_gz_keeps_values_exactly(tmp_path, edge_schema):
    query = edge_query(edge_schema)
    code, output = run_export(tmp_path, "edge.csv.gz", query)

    assert code == 0
    with gzip.open(output, "rt", encoding="utf-8-sig", newline="") as handle:
        raw = handle.read()
    rows = list(csv.reader(raw.splitlines(keepends=True)))
    assert rows[0][:3] == ["patient_id", "patient_name", "description"]
    assert len(rows) == 6
    assert rows[1][2] == LONG_DESCRIPTION
    assert rows[3][1] == "Müller 患者 😀"
    assert rows[4][2] == "bell\x07char"
    assert rows[1][4] == "2024-01-20 05:30:00+00:00" or rows[1][4].startswith("2024-01-20")
    assert rows[1][6] == "12345678901234567890.123"
    assert '"2","=SUM(A1:A2)","",,' in raw


def test_text_longer_than_an_excel_cell_is_reported_not_truncated(tmp_path, caplog):
    query = "SELECT 7 AS patient_id, repeat('x', 40000) AS description"
    code, output = run_export(tmp_path, "too_long.xlsx", query)

    assert code == 1
    assert "column description: 1 value(s)" in caplog.text
    assert "patient_id = 7" in caplog.text
    assert not output.parent.exists() or not list(output.parent.iterdir())


def test_rows_are_split_across_sheets_and_workbooks(tmp_path):
    query = ex.DEFAULT_QUERY_FILE.read_text(encoding="utf-8")
    total = query_rows(f"SELECT count(*) FROM ({query}) q")[0][0]
    code, output = run_export(
        tmp_path, "split.xlsx", None, "--rows-per-sheet", "1000", "--sheets-per-workbook", "2"
    )

    assert code == 0
    files = sorted(output.parent.glob("*.xlsx"))
    expected_sheets = -(-total // 1000)
    assert len(files) == -(-expected_sheets // 2)
    rows = [row for path in files for row in data_rows(path)]
    assert [row[0] for row in rows] == [row[0] for row in query_rows(query)]
    first = load_workbook(files[0], read_only=True)
    assert first.sheetnames == ["Data 1", "Data 2"][: min(2, expected_sheets)]
    first.close()


def test_limit_exports_the_first_rows(tmp_path):
    code, output = run_export(tmp_path, "first10.csv", None, "--limit", "10")

    assert code == 0
    with open(output, encoding="utf-8-sig", newline="") as handle:
        assert len(list(csv.reader(handle))) == 11


def test_existing_output_is_not_overwritten(tmp_path):
    output = tmp_path / "out" / "keep.csv"
    output.parent.mkdir()
    output.write_text("keep me")

    code, _ = run_export(tmp_path, "keep.csv", None, "--limit", "5")
    assert code == 1
    assert output.read_text() == "keep me"

    code, _ = run_export(tmp_path, "keep.csv", None, "--limit", "5", "--overwrite")
    assert code == 0
    assert output.read_text(encoding="utf-8-sig") != "keep me"


def test_missing_password_is_explained(tmp_path, monkeypatch, caplog):
    monkeypatch.delenv("PGPASSWORD")
    code, _ = run_export(tmp_path, "x.xlsx")
    assert code == 1
    assert "$env:PGPASSWORD" in caplog.text


SHOWCASE_DB = os.environ.get("EXPORT_SHOWCASE_DB", "loadtest")


def test_fifty_thousand_rows_with_long_descriptions(tmp_path):
    database = {**DB, "dbname": SHOWCASE_DB}
    try:
        with psycopg2.connect(password=os.environ["PGPASSWORD"], **database) as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    "SELECT count(*) FROM information_schema.columns "
                    "WHERE table_name = 'showcase_records' AND column_name = 'description'"
                )
                exists = cursor.fetchone()[0] == 1
    except psycopg2.OperationalError:
        exists = False
    if not exists:
        pytest.skip(f"showcase_records with a description column is not in {SHOWCASE_DB}")
    query = "SELECT id, name, email, city, amount, description FROM showcase_records ORDER BY id"
    code, output = run_export(tmp_path, "showcase.xlsx", query, "--dbname", SHOWCASE_DB)

    assert code == 0
    rows = data_rows(output)
    assert len(rows) >= 50_000
    with psycopg2.connect(password=os.environ["PGPASSWORD"], **database) as connection:
        with connection.cursor(name="check") as cursor:
            cursor.execute(query)
            number = 0
            while batch := cursor.fetchmany(5000):
                for source in batch:
                    read = tuple(rows[number]) + (None,) * (len(source) - len(rows[number]))
                    number += 1
                    for column, (r, s) in enumerate(zip(read, source)):
                        assert_same(r, excel_expected(s), f"row {number}, column {column + 1}")
    assert number == len(rows)
    print(f"\n{len(rows):,} showcase rows exported and compared value by value")
