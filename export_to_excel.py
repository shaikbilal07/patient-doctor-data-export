import argparse
import csv
import datetime as dt
import gzip
import json
import logging
import math
import os
import re
import shutil
import sys
import tempfile
import time
from copy import copy
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import psycopg2
import psycopg2.extensions
from psycopg2 import sql
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ERROR_CODES, ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

log = logging.getLogger("pg_export")

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_QUERY_FILE = SCRIPT_DIR / "patient_doctor_query.sql"

EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_DATA_ROWS = EXCEL_MAX_ROWS - 1
EXCEL_MAX_COLUMNS = 16_384
EXCEL_MAX_CELL_CHARS = 32_767
EXCEL_EXACT_DIGITS = 15
EXCEL_FIRST_DATE = dt.date(1900, 3, 1)

INT_TYPES = {20, 21, 23}
FLOAT_TYPES = {700, 701}
NUMERIC_TYPE = 1700
BOOL_TYPE = 16
DATE_TYPE = 1082
TIME_TYPE = 1083
TIMESTAMP_TYPES = {1114, 1184}
TIMESTAMPTZ_TYPE = 1184

NUMBER_FORMATS = {
    "int": "0",
    "date": "yyyy-mm-dd",
    "datetime": "yyyy-mm-dd hh:mm:ss",
    "time": "hh:mm:ss",
}
MIN_WIDTHS = {"date": 12, "datetime": 20, "time": 10}
MIN_COLUMN_WIDTH = 8
LINE_HEIGHT_POINTS = 15
WIDTH_SAMPLE_ROWS = 1_000

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
HEADER_BORDER = Border(bottom=Side(style="thin", color="000000"))
TOP_ALIGNMENT = Alignment(vertical="top")
WRAP_ALIGNMENT = Alignment(vertical="top", wrap_text=True)


class ExportError(Exception):
    pass


@dataclass
class Column:
    name: str
    type_code: int
    scale: int | None
    header: str = ""
    kind: str = "text"
    width: float = MIN_COLUMN_WIDTH
    wrap: bool = False
    number_format: str | None = None


@dataclass
class ExportResult:
    rows: int = 0
    sheets: int = 0
    files: list[Path] = field(default_factory=list)
    columns_with_values: set[int] = field(default_factory=set)
    escaped_values: int = 0
    formula_like_values: int = 0


class Progress:

    def __init__(self, interval: float, expected: int | None):
        self.interval = interval
        self.expected = expected
        self.started = self.last = time.monotonic()

    def update(self, rows: int, force: bool = False) -> None:
        now = time.monotonic()
        if not force and now - self.last < self.interval:
            return
        self.last = now
        elapsed = now - self.started
        speed = rows / elapsed if elapsed > 0 else 0
        message = f"Progress: {rows:,} rows"
        if self.expected:
            message += f" of {self.expected:,} ({rows / self.expected:.1%})"
        message += f", elapsed {format_duration(elapsed)}, {speed:,.0f} rows/s"
        if self.expected and speed > 0 and rows < self.expected:
            message += f", about {format_duration((self.expected - rows) / speed)} left"
        log.info(message)


def format_duration(seconds: float) -> str:
    seconds = int(round(seconds))
    hours, rest = divmod(seconds, 3600)
    minutes, secs = divmod(rest, 60)
    if hours:
        return f"{hours}h {minutes:02d}m {secs:02d}s"
    if minutes:
        return f"{minutes}m {secs:02d}s"
    return f"{secs}s"


def format_bytes(size: float) -> str:
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if size < 1024 or unit == "TB":
            return f"{size:,.1f} {unit}" if unit != "B" else f"{int(size)} B"
        size /= 1024
    return f"{size:,.1f} TB"


def pretty_header(name: str) -> str:
    if name != name.lower() or " " in name:
        return name
    words = [w for w in name.split("_") if w]
    return " ".join("ID" if w == "id" else w.capitalize() for w in words) or name


def excel_text_length(text: str) -> int:
    return len(text.encode("utf-16-le")) // 2


def output_format(path: Path) -> str:
    name = path.name.lower()
    for suffix in (".csv.gz", ".csv", ".xlsx"):
        if name.endswith(suffix) and len(name) > len(suffix):
            return suffix[1:]
    raise ExportError(f"{path.name}: the output file must end with .xlsx, .csv or .csv.gz")


def part_path(path: Path) -> Path:
    return path.with_name(path.name + ".part")


def connect(args) -> "psycopg2.extensions.connection":
    password = os.environ.get(args.password_env)
    if not password:
        raise ExportError(
            f"The database password is not set. In PowerShell run:\n"
            f'    $env:{args.password_env} = "your-password"\n'
            f"and then start the export again in the same window."
        )
    log.info(
        "Connecting to %s@%s:%s/%s", args.user, args.host, args.port, args.dbname
    )
    connection = psycopg2.connect(
        host=args.host,
        port=args.port,
        dbname=args.dbname,
        user=args.user,
        password=password,
        connect_timeout=15,
        application_name="pg_export",
    )
    connection.set_session(
        isolation_level=psycopg2.extensions.ISOLATION_LEVEL_REPEATABLE_READ,
        readonly=True,
    )
    return connection


def read_query(path: Path, limit: int | None) -> str:
    if not path.is_file():
        raise ExportError(f"The query file {path} does not exist")
    query = path.read_text(encoding="utf-8-sig").strip().rstrip(";").strip()
    if not query:
        raise ExportError(f"The query file {path} is empty")
    if ";" in re.sub(r"--[^\n]*", "", query):
        raise ExportError(
            f"The query file {path} must contain exactly one SELECT statement"
        )
    if limit is not None:
        query = f"{query}\nLIMIT {int(limit)}"
    return query


def column_kind(type_code: int) -> str:
    if type_code in INT_TYPES:
        return "int"
    if type_code in FLOAT_TYPES:
        return "float"
    if type_code == NUMERIC_TYPE:
        return "numeric"
    if type_code == BOOL_TYPE:
        return "bool"
    if type_code == DATE_TYPE:
        return "date"
    if type_code in TIMESTAMP_TYPES:
        return "datetime"
    if type_code == TIME_TYPE:
        return "time"
    return "text"


def describe_columns(connection, query: str) -> list[Column]:
    statement = sql.SQL("SELECT * FROM ({}) AS q LIMIT 0").format(sql.SQL(query))
    with connection.cursor() as cursor:
        cursor.execute(statement)
        description = cursor.description
    columns = []
    for item in description:
        column = Column(name=item.name, type_code=item.type_code, scale=item.scale)
        column.kind = column_kind(item.type_code)
        column.header = pretty_header(item.name)
        if item.type_code == TIMESTAMPTZ_TYPE:
            column.header += " (UTC)"
        columns.append(column)

    names = [c.name for c in columns]
    duplicates = sorted({n for n in names if names.count(n) > 1})
    if duplicates:
        raise ExportError(
            "The query returns more than one column named "
            + ", ".join(duplicates)
            + ". Give each column a unique name with AS."
        )
    return columns


def estimate_result(connection, query: str) -> tuple[int, int]:
    with connection.cursor() as cursor:
        cursor.execute(sql.SQL("EXPLAIN (FORMAT JSON) {}").format(sql.SQL(query)))
        plan = cursor.fetchone()[0]
    if isinstance(plan, str):
        plan = json.loads(plan)
    top = plan[0]["Plan"]
    return int(top["Plan Rows"]), int(top["Plan Width"])


def may_hold_long_text(column: Column) -> bool:
    return column.kind == "text"


def count_rows(connection, query: str, columns: list[Column], check_cells: bool):
    text_columns = [c for c in columns if check_cells and may_hold_long_text(c)]
    parts = [sql.SQL("count(*)")]
    for column in text_columns:
        parts.append(
            sql.SQL("count(*) FILTER (WHERE length(q.{}::text) > {})").format(
                sql.Identifier(column.name), sql.Literal(EXCEL_MAX_CELL_CHARS)
            )
        )
    statement = sql.SQL("SELECT {} FROM ({}) AS q").format(
        sql.SQL(", ").join(parts), sql.SQL(query)
    )
    with connection.cursor() as cursor:
        cursor.execute(statement)
        result = cursor.fetchone()
    too_long = {c.name: n for c, n in zip(text_columns, result[1:]) if n}
    return result[0], too_long


def report_too_long_values(connection, query: str, columns, too_long: dict) -> None:
    key = columns[0]
    lines = []
    for name, count in too_long.items():
        statement = sql.SQL(
            "SELECT q.{key} FROM ({query}) AS q "
            "WHERE length(q.{col}::text) > {limit} LIMIT 5"
        ).format(
            key=sql.Identifier(key.name),
            query=sql.SQL(query),
            col=sql.Identifier(name),
            limit=sql.Literal(EXCEL_MAX_CELL_CHARS),
        )
        with connection.cursor() as cursor:
            cursor.execute(statement)
            examples = ", ".join(str(row[0]) for row in cursor.fetchall())
        lines.append(
            f"  column {name}: {count:,} value(s), for example in rows where "
            f"{key.name} = {examples}"
        )
    raise ExportError(
        f"Some values are longer than Excel's limit of {EXCEL_MAX_CELL_CHARS:,} "
        f"characters per cell:\n" + "\n".join(lines) + "\n"
        "Nothing was written, because Excel would cut these values short. "
        "Export to .csv or .csv.gz instead, which keeps every value complete."
    )


def check_excel_estimate(estimated_rows: int, args) -> None:
    if estimated_rows > args.max_xlsx_rows * 10:
        raise ExportError(
            f"PostgreSQL estimates the query returns about {estimated_rows:,} rows. "
            f"That is far more than this program writes to Excel (--max-xlsx-rows "
            f"{args.max_xlsx_rows:,}). {EXCEL_ADVICE}"
        )


def check_excel_capacity(columns: list[Column], total_rows: int, args) -> None:
    if len(columns) > EXCEL_MAX_COLUMNS:
        raise ExportError(
            f"The query returns {len(columns):,} columns, but an Excel worksheet "
            f"holds at most {EXCEL_MAX_COLUMNS:,}."
        )
    if total_rows > args.max_xlsx_rows:
        raise ExportError(
            f"The query returns {total_rows:,} rows, more than --max-xlsx-rows "
            f"{args.max_xlsx_rows:,}. {EXCEL_ADVICE}"
        )


EXCEL_ADVICE = (
    "One worksheet holds at most 1,048,575 data rows (1,048,576 including the "
    "header), and Excel becomes very slow or fails to open files with millions of "
    "rows. For a result this large, export to .csv.gz: it has no row limit, keeps "
    "every value exactly, and is several times smaller."
)


def check_disk_space(fmt, estimated_rows, estimated_width, columns, outputs) -> None:
    raw = estimated_rows * (estimated_width + 2 * len(columns))
    output_dir = outputs[0].parent
    needs: dict[Path, float] = {}
    if fmt == "csv":
        needs[output_dir] = raw * 1.3
    elif fmt == "csv.gz":
        needs[output_dir] = raw * 0.5
    else:
        xml = raw + estimated_rows * len(columns) * 60
        needs[Path(tempfile.gettempdir())] = xml
        needs[output_dir] = raw * 0.7

    by_drive: dict[str, tuple[Path, float]] = {}
    for path, need in needs.items():
        anchor = path.resolve().anchor.upper()
        known = by_drive.get(anchor, (path, 0.0))
        by_drive[anchor] = (known[0], known[1] + need)

    for anchor, (path, need) in by_drive.items():
        free = shutil.disk_usage(path).free
        log.info(
            "Disk %s: about %s needed (estimate), %s free",
            anchor,
            format_bytes(need),
            format_bytes(free),
        )
        if need * 1.1 > free:
            raise ExportError(
                f"Not enough free disk space on {anchor}: the export needs about "
                f"{format_bytes(need)} (estimated from PostgreSQL statistics) but only "
                f"{format_bytes(free)} is free. Free up space, choose another "
                f"--output drive, or use .csv.gz. Use --skip-disk-check to try anyway."
            )


class TextCleaner:

    def __init__(self):
        self.escaped_values = 0

    def clean(self, text: str) -> str:
        if ILLEGAL_CHARACTERS_RE.search(text):
            self.escaped_values += 1
            text = ILLEGAL_CHARACTERS_RE.sub(lambda m: f"\\x{ord(m.group()):02x}", text)
        return text


def plain_value(value):
    if isinstance(value, (memoryview, bytes, bytearray)):
        return "\\x" + bytes(value).hex()
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, dt.timedelta):
        return str(value)
    return value


def to_excel(value, cleaner: TextCleaner):
    value = plain_value(value)
    if isinstance(value, bool):
        return value, False
    if isinstance(value, int):
        if len(str(abs(value))) > EXCEL_EXACT_DIGITS:
            return str(value), True
        return value, False
    if isinstance(value, float):
        return (value, False) if math.isfinite(value) else (str(value), True)
    if isinstance(value, Decimal):
        if value.is_finite() and len(value.as_tuple().digits) <= EXCEL_EXACT_DIGITS:
            return value, False
        return str(value), True
    if isinstance(value, dt.datetime):
        if value.tzinfo is not None:
            value = value.astimezone(dt.timezone.utc).replace(tzinfo=None)
        if value.date() < EXCEL_FIRST_DATE:
            return value.isoformat(" "), True
        return value, False
    if isinstance(value, dt.date):
        return (value, False) if value >= EXCEL_FIRST_DATE else (value.isoformat(), True)
    if isinstance(value, dt.time):
        return (value, False) if value.tzinfo is None else (value.isoformat(), True)
    return cleaner.clean(str(value)), True


def to_csv(value):
    value = plain_value(value)
    if value is None or isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, dt.datetime):
        return value.isoformat(" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def plan_layout(columns: list[Column], sample_rows, max_width: int) -> None:
    for index, column in enumerate(columns):
        lengths = [
            len(str(plain_value(row[index]))) for row in sample_rows if row[index] is not None
        ]
        longest = max(lengths, default=0)
        header_need = min(len(column.header), max_width)
        width = max(longest, header_need, MIN_COLUMN_WIDTH, MIN_WIDTHS.get(column.kind, 0))
        column.wrap = column.kind == "text" and longest > max_width
        column.width = max_width if column.wrap else min(width + 2, max_width)
        if column.kind == "numeric" and column.scale is not None:
            column.number_format = "#,##0" + ("." + "0" * column.scale if column.scale else "")
        else:
            column.number_format = NUMBER_FORMATS.get(column.kind)


def wrapped_lines(text: str, width: float) -> int:
    per_line = max(int(width * 1.1), 1)
    return sum(max(1, math.ceil(len(part) / per_line)) for part in text.split("\n"))


class ExcelWriter:

    def __init__(self, columns, final_paths, args, result: ExportResult):
        self.columns = columns
        self.final_paths = final_paths
        self.rows_per_sheet = args.rows_per_sheet
        self.sheets_per_workbook = args.sheets_per_workbook
        self.max_row_lines = args.max_row_lines
        self.total_sheets = None
        self.result = result
        self.cleaner = TextCleaner()
        self.workbook = None
        self.sheet = None
        self.styles = {}
        self.workbook_index = 0
        self.sheets_in_workbook = 0
        self.rows_in_sheet = 0
        self.last_column = get_column_letter(len(columns))

    def plan(self, total_rows: int) -> None:
        self.total_sheets = max(1, math.ceil(total_rows / self.rows_per_sheet))

    def write_rows(self, rows) -> None:
        for row in rows:
            if self.sheet is None or self.rows_in_sheet == self.rows_per_sheet:
                self._next_sheet()
            self._write_row(row)

    def _next_sheet(self) -> None:
        if self.sheet is not None:
            self._finish_sheet()
        if self.workbook is None or self.sheets_in_workbook == self.sheets_per_workbook:
            self._save_workbook()
            self.workbook = Workbook(write_only=True)
            self.sheets_in_workbook = 0
            self.styles = {}

        self.result.sheets += 1
        self.sheets_in_workbook += 1
        title = "Data" if self.total_sheets == 1 else f"Data {self.result.sheets}"
        sheet = self.workbook.create_sheet(title)
        for index, column in enumerate(self.columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = column.width
        sheet.freeze_panes = "A2"
        header_lines = max(wrapped_lines(c.header, c.width) for c in self.columns)
        sheet.row_dimensions[1].height = LINE_HEIGHT_POINTS * min(header_lines, 3) + 6
        header = []
        for column in self.columns:
            cell = WriteOnlyCell(sheet, column.header)
            cell.data_type = "s"
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT
            cell.border = HEADER_BORDER
            header.append(cell)
        sheet.append(header)
        self.sheet = sheet
        self.rows_in_sheet = 0

    def _write_row(self, row) -> None:
        sheet_row = self.rows_in_sheet + 2
        cells = []
        lines = 1
        for index, (column, raw) in enumerate(zip(self.columns, row)):
            if raw is None:
                cells.append(None)
                continue
            self.result.columns_with_values.add(index)
            value, is_text = to_excel(raw, self.cleaner)
            cell = WriteOnlyCell(self.sheet, value)
            if is_text:
                if excel_text_length(value) > EXCEL_MAX_CELL_CHARS:
                    raise ExportError(
                        f"Row {self.result.rows + 1:,}, column {column.name}: the value "
                        f"has more than {EXCEL_MAX_CELL_CHARS:,} characters, which does "
                        "not fit in an Excel cell. Nothing was saved; export to .csv.gz "
                        "to keep the full value."
                    )
                if value.startswith("=") or value in ERROR_CODES:
                    self.result.formula_like_values += 1
                cell.data_type = "s"
                if column.wrap or "\n" in value:
                    cell._style = self._style(WRAP_ALIGNMENT, None)
                    lines = max(lines, wrapped_lines(value, column.width))
                else:
                    cell._style = self._style(TOP_ALIGNMENT, None)
            elif column.number_format and not isinstance(value, bool):
                cell._style = self._style(TOP_ALIGNMENT, column.number_format)
            else:
                cell._style = self._style(TOP_ALIGNMENT, None)
            cells.append(cell)

        if lines > 1 and self.max_row_lines > 0:
            height = LINE_HEIGHT_POINTS * min(lines, self.max_row_lines)
            self.sheet.row_dimensions[sheet_row].height = height
            self.sheet.append(cells)
            del self.sheet.row_dimensions[sheet_row]
        else:
            self.sheet.append(cells)
        self.rows_in_sheet += 1
        self.result.rows += 1

    def _style(self, alignment: Alignment, number_format: str | None):
        key = (id(alignment), number_format)
        style = self.styles.get(key)
        if style is None:
            template = WriteOnlyCell(self.sheet)
            template.alignment = alignment
            if number_format:
                template.number_format = number_format
            style = self.styles[key] = template._style
        return copy(style)

    def _finish_sheet(self) -> None:
        last_row = self.rows_in_sheet + 1
        self.sheet.auto_filter.ref = f"A1:{self.last_column}{last_row}"

    def _save_workbook(self) -> None:
        if self.workbook is None:
            return
        if self.workbook_index >= len(self.final_paths):
            raise ExportError(
                "The query returned more rows than it did when they were counted"
            )
        path = part_path(self.final_paths[self.workbook_index])
        log.info("Saving workbook %s", path.name)
        self.workbook.save(path)
        self.result.files.append(path)
        self.workbook_index += 1
        self.workbook = None

    def finish(self) -> None:
        if self.sheet is not None:
            self._finish_sheet()
        self._save_workbook()
        self.result.escaped_values = self.cleaner.escaped_values

    def abort(self) -> None:
        if self.workbook is not None:
            try:
                for sheet in self.workbook.worksheets:
                    if not sheet.closed:
                        sheet.close()
            except Exception:
                pass
            self.workbook = None


def xlsx_paths(output: Path, total_rows: int, args) -> list[Path]:
    sheets = max(1, math.ceil(total_rows / args.rows_per_sheet))
    workbooks = math.ceil(sheets / args.sheets_per_workbook)
    if workbooks == 1:
        return [output]
    return [
        output.with_name(f"{output.stem}_part{n:03d}{output.suffix}")
        for n in range(1, workbooks + 1)
    ]


def stream_batches(connection, query: str, batch_size: int):
    cursor = connection.cursor(name="pg_export_cursor")
    try:
        cursor.execute(query)
        while True:
            batch = cursor.fetchmany(batch_size)
            if not batch:
                return
            yield batch
    finally:
        try:
            cursor.close()
        except psycopg2.Error:
            pass


def export_xlsx(connection, query, columns, paths, expected, args, result, progress):
    writer = ExcelWriter(columns, paths, args, result)
    writer.plan(expected)
    try:
        for batch in stream_batches(connection, query, args.batch_size):
            if result.rows == 0:
                plan_layout(columns, batch[:WIDTH_SAMPLE_ROWS], args.max_column_width)
                wrapped = [c.name for c in columns if c.wrap]
                if wrapped:
                    log.info("Long text is wrapped in column(s): %s", ", ".join(wrapped))
            writer.write_rows(batch)
            progress.update(result.rows)
        writer.finish()
    except BaseException:
        writer.abort()
        raise


def export_csv(connection, query, columns, path, args, result, progress):
    target = part_path(path)
    quoting = getattr(csv, "QUOTE_NOTNULL", csv.QUOTE_MINIMAL)
    if path.name.lower().endswith(".gz"):
        handle = gzip.open(
            target, "wt", encoding="utf-8-sig", newline="", compresslevel=args.gzip_level
        )
    else:
        handle = open(target, "w", encoding="utf-8-sig", newline="", buffering=1 << 20)
    result.files.append(target)
    with handle:
        writer = csv.writer(handle, quoting=quoting, lineterminator="\r\n")
        writer.writerow([c.name for c in columns])
        for batch in stream_batches(connection, query, args.batch_size):
            out = []
            for row in batch:
                values = [to_csv(v) for v in row]
                for index, value in enumerate(values):
                    if value is None:
                        continue
                    result.columns_with_values.add(index)
                    if value[:1] in ("=", "+", "-", "@") and not _is_number(value):
                        result.formula_like_values += 1
                        if args.csv_escape_formulas:
                            values[index] = "'" + value
                out.append(values)
            writer.writerows(out)
            result.rows += len(batch)
            progress.update(result.rows)


def _is_number(text: str) -> bool:
    try:
        float(text)
        return True
    except ValueError:
        return False


def run(args) -> int:
    started = time.monotonic()
    output = args.output.resolve()
    fmt = output_format(output)
    query = read_query(args.query_file, args.limit)
    result = ExportResult()

    if not Path(output.anchor).exists():
        raise ExportError(f"The drive {output.anchor} does not exist")
    output.parent.mkdir(parents=True, exist_ok=True)

    connection = connect(args)
    try:
        columns = describe_columns(connection, query)
        log.info("The query returns %d columns: %s", len(columns), ", ".join(c.name for c in columns))
        est_rows, est_width = estimate_result(connection, query)
        log.info(
            "PostgreSQL estimates about %s rows of about %s bytes each",
            f"{est_rows:,}",
            est_width,
        )
        if fmt == "xlsx":
            check_excel_estimate(est_rows, args)
        if not args.skip_disk_check:
            check_disk_space(fmt, est_rows, est_width, columns, [output])

        expected = None
        if fmt == "xlsx" or not args.skip_count:
            log.info("Counting the rows of the query (this reads the whole result once)")
            expected, too_long = count_rows(connection, query, columns, fmt == "xlsx")
            log.info("The query returns %s rows", f"{expected:,}")
            if too_long:
                report_too_long_values(connection, query, columns, too_long)
            if expected == 0:
                log.error("The query returned no rows, so no file was written")
                return 1

        if fmt == "xlsx":
            check_excel_capacity(columns, expected, args)
            paths = xlsx_paths(output, expected, args)
        else:
            paths = [output]
        existing = [p for p in paths if p.exists()]
        if existing and not args.overwrite:
            raise ExportError(
                f"{existing[0]} already exists. Use --overwrite to replace it."
            )
        for path in paths:
            part_path(path).unlink(missing_ok=True)
        if len(paths) > 1:
            log.info("The result is split into %d workbooks", len(paths))

        progress = Progress(args.progress_seconds, expected)
        try:
            if fmt == "xlsx":
                export_xlsx(connection, query, columns, paths, expected, args, result, progress)
            else:
                export_csv(connection, query, columns, output, args, result, progress)
        except BaseException:
            for path in paths:
                part_path(path).unlink(missing_ok=True)
            raise
        progress.update(result.rows, force=True)
    finally:
        connection.close()

    if result.rows == 0:
        for path in paths:
            part_path(path).unlink(missing_ok=True)
        log.error("The query returned no rows, so no file was written")
        return 1
    if expected is not None and expected != result.rows:
        log.error(
            "Verification FAILED: the query has %s rows but %s were written. "
            "The unverified file(s) were kept as: %s",
            f"{expected:,}",
            f"{result.rows:,}",
            ", ".join(str(p) for p in result.files),
        )
        return 1

    for path in paths:
        os.replace(part_path(path), path)
    print_summary(args, fmt, paths, columns, result, expected, started)
    return 0


def print_summary(args, fmt, paths, columns, result, expected, started) -> None:
    seconds = time.monotonic() - started
    size = sum(p.stat().st_size for p in paths)
    lines = [
        "Export finished",
        f"  Output file(s)    : {', '.join(str(p) for p in paths)}",
        f"  Format            : {fmt}",
        f"  Total file size   : {format_bytes(size)}",
        f"  Rows exported     : {result.rows:,}",
        f"  Columns           : {len(columns)}",
    ]
    if result.sheets:
        lines.append(f"  Worksheets        : {result.sheets}")
    if expected is not None:
        lines.append(f"  Row count check   : query {expected:,} = written {result.rows:,} (OK)")
    else:
        lines.append("  Row count check   : skipped (--skip-count)")
    lines += [
        f"  Elapsed time      : {format_duration(seconds)}",
        f"  Throughput        : {result.rows / seconds if seconds else 0:,.0f} rows/s, "
        f"{format_bytes(size / seconds if seconds else 0)}/s written",
    ]
    log.info("\n".join(lines))

    empty = [c.name for i, c in enumerate(columns) if i not in result.columns_with_values]
    if empty:
        log.warning("These columns are NULL in every row: %s", ", ".join(empty))
    if result.escaped_values:
        log.warning(
            "%s text value(s) contained control characters that Excel cannot store; "
            "they were written as visible escapes such as \\x07. Use .csv for the exact "
            "bytes.",
            f"{result.escaped_values:,}",
        )
    if result.formula_like_values:
        if fmt == "xlsx":
            log.info(
                "%s value(s) start with '=' or look like Excel errors; they were stored "
                "as plain text, not formulas.",
                f"{result.formula_like_values:,}",
            )
        elif args.csv_escape_formulas:
            log.info(
                "%s value(s) that start with = + - @ were prefixed with ' so Excel "
                "does not run them as formulas.",
                f"{result.formula_like_values:,}",
            )
        else:
            log.warning(
                "%s value(s) start with = + - or @. Excel may treat them as formulas "
                "when the CSV is opened by double-clicking. Use --csv-escape-formulas "
                "if the file is meant for Excel users.",
                f"{result.formula_like_values:,}",
            )


def positive_int(text: str) -> int:
    value = int(text)
    if value < 1:
        raise argparse.ArgumentTypeError("must be 1 or more")
    return value


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description="Export a PostgreSQL query result to .xlsx, .csv or .csv.gz.",
        epilog="The password is read from the environment variable named by "
        "--password-env (default PGPASSWORD); it is never accepted on the command line.",
    )
    parser.add_argument("--output", type=Path, required=True,
                        help="output file; the extension .xlsx, .csv or .csv.gz selects the format")
    parser.add_argument("--query-file", type=Path, default=DEFAULT_QUERY_FILE,
                        help="file with one SELECT statement, without LIMIT "
                        "(default: patient_doctor_query.sql)")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=5432)
    parser.add_argument("--dbname", default="hospital_export_test")
    parser.add_argument("--user", default="postgres")
    parser.add_argument("--password-env", default="PGPASSWORD",
                        help="environment variable holding the password (default: PGPASSWORD)")
    parser.add_argument("--batch-size", type=positive_int, default=10_000,
                        help="rows fetched from the server per batch (default: 10000)")
    parser.add_argument("--limit", type=positive_int,
                        help="export only the first N rows, for testing")
    parser.add_argument("--rows-per-sheet", type=positive_int, default=EXCEL_MAX_DATA_ROWS,
                        help="data rows per worksheet, at most 1048575 (default: 1048575)")
    parser.add_argument("--sheets-per-workbook", type=positive_int, default=1,
                        help="worksheets per .xlsx file before a new file is started "
                        "(default: 1, which keeps each file easy for Excel to open)")
    parser.add_argument("--max-xlsx-rows", type=positive_int, default=5_000_000,
                        help="refuse .xlsx exports with more rows than this (default: 5000000)")
    parser.add_argument("--max-column-width", type=positive_int, default=60,
                        help="widest column in Excel character units; longer text wraps "
                        "(default: 60)")
    parser.add_argument("--max-row-lines", type=int, default=10,
                        help="tallest row, in lines of wrapped text; the full text stays in "
                        "the cell. 0 lets Excel decide (default: 10)")
    parser.add_argument("--gzip-level", type=int, choices=range(1, 10), default=1,
                        help="compression level for .csv.gz; 1 is fastest (default: 1)")
    parser.add_argument("--csv-escape-formulas", action="store_true",
                        help="prefix CSV text starting with = + - @ with ' so Excel shows "
                        "it as text (changes those values)")
    parser.add_argument("--skip-count", action="store_true",
                        help="CSV only: do not count the rows first (saves one full read "
                        "of the data, but the row count is not verified)")
    parser.add_argument("--skip-disk-check", action="store_true",
                        help="do not refuse to start when the disk space estimate is too low")
    parser.add_argument("--overwrite", action="store_true",
                        help="replace existing output files")
    parser.add_argument("--progress-seconds", type=float, default=10,
                        help="seconds between progress messages (default: 10)")
    parser.add_argument("--log-file", type=Path, help="also write the log to this file")
    args = parser.parse_args(argv)
    if args.rows_per_sheet > EXCEL_MAX_DATA_ROWS:
        parser.error(f"--rows-per-sheet cannot be more than {EXCEL_MAX_DATA_ROWS}")
    try:
        output_format(args.output)
    except ExportError as exc:
        parser.error(str(exc))
    return args


def configure_logging(log_file: Path | None) -> None:
    for handler in log.handlers[:]:
        log.removeHandler(handler)
        handler.close()
    handlers = [logging.StreamHandler(sys.stdout)]
    if log_file:
        handlers.append(logging.FileHandler(log_file, encoding="utf-8"))
    formatter = logging.Formatter("%(asctime)s %(levelname)-7s %(message)s", "%H:%M:%S")
    for handler in handlers:
        handler.setFormatter(formatter)
        log.addHandler(handler)
    log.setLevel(logging.INFO)


def main(argv=None) -> int:
    args = parse_args(argv)
    configure_logging(args.log_file)
    try:
        return run(args)
    except KeyboardInterrupt:
        log.warning(
            "Interrupted. The partial output was deleted. This export cannot resume; "
            "run it again from the start."
        )
        return 130
    except ExportError as exc:
        log.error("%s", exc)
        return 1
    except psycopg2.Error as exc:
        message = (exc.pgerror or str(exc)).strip()
        log.error("Database error: %s", message)
        return 1
    except OSError as exc:
        log.error("File error: %s", exc)
        return 1


if __name__ == "__main__":
    sys.exit(main())
